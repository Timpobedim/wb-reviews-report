from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from typing import Any

import pytest
from conftest import FakeWb
from openpyxl import load_workbook

from wb_reviews import demo
from wb_reviews.cli import main

DRESS, SHOES, MUG = demo.SCENARIOS
NOW = "2026-09-10T12:00"


def run(wb_args: list[str], tmp_path: Path, out: Path) -> int:
    return main(["report", *wb_args, "-o", str(out), "--cache", str(tmp_path / "cache"), "--pause", "0", "--now", NOW])


def cells(path: Path) -> str:
    workbook = load_workbook(path)
    return "\n".join(
        str(value)
        for sheet in workbook
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )


def summary_rows(path: Path) -> dict[Any, dict[str, Any]]:
    sheet = load_workbook(path)["Сводка"]
    header = [cell.value for cell in sheet[3]]
    return {row[0]: dict(zip(header, row, strict=True)) for row in sheet.iter_rows(min_row=4, values_only=True)}


def test_demo_report(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "демо.xlsx"
    assert main(["demo", "-o", str(out)]) == 0
    workbook = load_workbook(out)
    assert workbook.sheetnames == ["Сводка", "Жалобы", "Оценки", "По месяцам", "Отзывы 1–3★"]
    rows = summary_rows(out)
    assert [rows[s.nm]["Тренд"] for s in demo.SCENARIOS] == [
        "негатив растёт",
        "без изменений",
        "негатив снижается",
    ]
    assert workbook["Сводка"]["H4"].number_format == "0%"
    assert len(workbook["По месяцам"]._charts) == 1
    complaints = workbook["Отзывы 1–3★"]
    ratings = [row[4] for row in complaints.iter_rows(min_row=4, values_only=True)]
    assert ratings
    assert max(ratings) <= 3
    assert len(ratings) <= 3 * 50
    printed = capsys.readouterr().out
    assert "Негатив растёт" in printed
    assert "Отчёт:" in printed


def test_report_has_no_personal_data_and_hides_contacts(wb: FakeWb, tmp_path: Path) -> None:
    wb.add(DRESS)
    wb.feedbacks[DRESS.imt]["feedbacks"][0].update(
        productValuation=1,
        text="Звоните +7 (912) 345-67-89 или пишите shop@example.com",
        excludedFromRating={"isExcluded": False, "reasons": []},
    )
    out = tmp_path / "отзывы.xlsx"
    assert run([str(DRESS.nm)], tmp_path, out) == 0
    text = cells(out)
    assert "Звоните [скрыто] или пишите [скрыто]" in text
    for secret in ("345-67-89", "shop@example.com", "Покупатель"):
        assert secret not in text


def test_report_lists_products_that_failed(wb: FakeWb, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    wb.add(SHOES)
    out = tmp_path / "отзывы.xlsx"
    assert run(["123456789", f"https://www.wildberries.ru/catalog/{SHOES.nm}/detail.aspx"], tmp_path, out) == 0
    failed = load_workbook(out)["Не удалось"]
    assert failed["A2"].value == "123456789"
    assert list(summary_rows(out)) == [SHOES.nm]
    assert "123456789" in capsys.readouterr().err


def test_nothing_collected_means_no_report(wb: FakeWb, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "отзывы.xlsx"
    assert run(["123456789"], tmp_path, out) == 2
    assert not out.exists()
    assert "Отчёт не построен" in capsys.readouterr().err


def test_products_from_file(wb: FakeWb, tmp_path: Path) -> None:
    wb.add(DRESS)
    wb.add(MUG)
    listing = tmp_path / "артикулы.txt"
    listing.write_text(
        f"# конкуренты\n{DRESS.nm}\n\nhttps://www.wildberries.ru/catalog/{MUG.nm}/detail.aspx\n{DRESS.nm}\n",
        encoding="utf-8",
    )
    out = tmp_path / "отзывы.xlsx"
    assert run(["-f", str(listing)], tmp_path, out) == 0
    assert list(summary_rows(out)) == [DRESS.nm, MUG.nm]


def test_own_theme_dictionary_replaces_the_default(wb: FakeWb, tmp_path: Path) -> None:
    wb.add(MUG)
    wb.feedbacks[MUG.imt]["feedbacks"][0].update(
        productValuation=1,
        text="Крышка не работает, клапан заедает",
        cons="",
        bables=None,
        excludedFromRating={"isExcluded": False, "reasons": []},
    )
    themes = tmp_path / "темы.json"
    themes.write_text(json.dumps({"Клапан и крышка": ["клапан", "крышк"]}, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "отзывы.xlsx"
    assert run([str(MUG.nm), "--themes", str(themes)], tmp_path, out) == 0
    workbook = load_workbook(out)
    rows = [row[0] for row in workbook["Жалобы"].iter_rows(min_row=4, values_only=True)]
    assert rows[0] == "Клапан и крышка"
    assert "Качество и брак" not in rows  # свой словарь заменяет стандартный
    newest = next(workbook["Отзывы 1–3★"].iter_rows(min_row=4, values_only=True))
    assert newest[5] == "Клапан и крышка"


def test_bad_theme_dictionary_is_explained(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    themes = tmp_path / "темы.json"
    themes.write_text('["маломер"]', encoding="utf-8")
    with pytest.raises(SystemExit) as stopped:
        main(["report", "13221580", "--themes", str(themes)])
    assert stopped.value.code == 2
    assert "нужен JSON-объект" in capsys.readouterr().err


def test_bad_input_is_explained(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as stopped:
        main(["report", "платье"])
    assert stopped.value.code == 2
    assert "Не похоже на артикул" in capsys.readouterr().err


def test_report_open_elsewhere_keeps_old_file(
    wb: FakeWb, tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    wb.add(DRESS)
    out = tmp_path / "отзывы.xlsx"
    out.write_bytes(b"old report")
    replace = os.replace

    def locked(source: str | Path, target: str | Path) -> None:
        if str(target).endswith(".xlsx"):
            raise PermissionError(13, "Отказано в доступе")
        replace(source, target)

    monkeypatch.setattr(os, "replace", locked)
    assert run([str(DRESS.nm)], tmp_path, out) == 1
    assert "Закройте его" in capsys.readouterr().err
    assert out.read_bytes() == b"old report"
    assert not list(tmp_path.glob("*.tmp"))


@pytest.mark.skipif(sys.platform != "win32", reason="блокировка открытого файла — поведение Windows")
def test_report_open_in_another_program_on_windows(
    wb: FakeWb, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    wb.add(DRESS)
    out = tmp_path / "отзывы.xlsx"
    out.write_bytes(b"old report")
    with out.open("rb"):  # как Excel: файл открыт, заменить его нельзя
        assert run([str(DRESS.nm)], tmp_path, out) == 1
    assert "Закройте его" in capsys.readouterr().err
    assert out.read_bytes() == b"old report"
