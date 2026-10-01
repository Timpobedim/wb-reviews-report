"""Отчёт в Excel: сводка, темы жалоб, оценки, негатив по месяцам и сами отзывы на 1–3★.

Файл пишется во временный рядом и только потом подменяет старый: если отчёт открыт в Excel, старый файл
остаётся целым, а пользователь получает понятное сообщение вместо трассировки.
"""

from __future__ import annotations

import os
from collections import Counter
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from wb_reviews.analysis import (
    BEFORE,
    DEFAULT_THEMES,
    MSK,
    NO_THEME,
    RECENT,
    SILENT,
    Summary,
    Themes,
    conclusions,
    monthly,
    themes_of,
)
from wb_reviews.models import ProductFeedback, mask_contacts

Item = tuple[ProductFeedback, Summary]
Column = tuple[str, float, str | None]  # заголовок, ширина, числовой формат

TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4B2A7B")
RISING_FILL = PatternFill("solid", fgColor="F4CCCC")
FALLING_FILL = PatternFill("solid", fgColor="D9EAD3")
TOP = Alignment(vertical="top", wrap_text=True)
HEAD = Alignment(vertical="center", wrap_text=True)
PERCENT, DATE, INTEGER = "0%", "DD.MM.YYYY", "#,##0"
QUOTES_PER_PRODUCT = 50
PRODUCTS = ("товар", "товара", "товаров")

SUMMARY_COLUMNS: tuple[Column, ...] = (
    ("Артикул WB", 12, "0"),
    ("Товар", 32, None),
    ("Предмет", 16, None),
    ("Рейтинг WB", 9, "0.0"),
    ("Отзывов всего", 10, INTEGER),
    ("В анализе", 10, INTEGER),
    ("Отзывы с", 11, DATE),
    ("1–2★ за 30 дней", 10, PERCENT),
    ("1–2★ за 90 дней до них", 12, PERCENT),
    ("Тренд", 16, None),
    ("Ответы продавца", 10, PERCENT),
    ("Ответы на 1–3★", 10, PERCENT),
    ("Медиана ответа, ч", 10, "0"),
    ("Маломерит", 10, PERCENT),
    ("Большемерит", 11, PERCENT),
    ("Главная жалоба", 22, None),
    ("Выводы", 80, None),
)
REVIEW_COLUMNS: tuple[Column, ...] = (
    ("Артикул WB", 12, "0"),
    ("Товар", 24, None),
    ("Вариант", 12, "0"),
    ("Дата", 11, DATE),
    ("Оценка", 8, "0"),
    ("Темы", 24, None),
    ("Теги WB", 22, None),
    ("Текст", 60, None),
    ("Недостатки", 40, None),
    ("Ответ продавца", 10, None),
)
TREND_FILLS = {"негатив растёт": RISING_FILL, "негатив снижается": FALLING_FILL}


class ReportError(Exception):
    pass


def plural(count: int, forms: tuple[str, str, str]) -> str:
    """1 товар, 2 товара, 5 товаров."""
    if count % 10 == 1 and count % 100 != 11:
        return forms[0]
    if 2 <= count % 10 <= 4 and not 12 <= count % 100 <= 14:
        return forms[1]
    return forms[2]


def _local(moment: datetime) -> datetime:
    """Excel не хранит часовой пояс: записываем московское время без пояса."""
    return moment.astimezone(MSK).replace(tzinfo=None)


def _short(name: str, limit: int = 30) -> str:
    return name if len(name) <= limit else name[: limit - 1] + "…"


def _label(summary: Summary) -> str:
    return f"{summary.card.nm}\n{_short(summary.card.name)}"


def _title(sheet: Worksheet, text: str) -> None:
    sheet.cell(1, 1, text).font = TITLE_FONT


def _table(
    sheet: Worksheet,
    top: int,
    columns: Sequence[Column],
    rows: Sequence[Sequence[Any]],
    *,
    autofilter: bool = True,
) -> None:
    for number, (title, width, _) in enumerate(columns, start=1):
        cell = sheet.cell(top, number, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEAD
        sheet.column_dimensions[get_column_letter(number)].width = width
    for offset, values in enumerate(rows, start=1):
        for number, (value, (_, _, number_format)) in enumerate(zip(values, columns, strict=True), start=1):
            cell = sheet.cell(top + offset, number, value)
            cell.alignment = TOP
            if number_format:
                cell.number_format = number_format
    if autofilter and rows:
        sheet.auto_filter.ref = f"A{top}:{get_column_letter(len(columns))}{top + len(rows)}"


def _summary(sheet: Worksheet, summaries: Sequence[Summary], now: datetime) -> None:
    count = len(summaries)
    _title(
        sheet,
        f"Отзывы на Wildberries: {count} {plural(count, PRODUCTS)}, данные на {_local(now):%d.%m.%Y %H:%M} МСК",
    )
    rows = [
        [
            s.card.nm,
            s.card.name,
            s.card.subject,
            s.rating,
            s.total,
            s.analyzed,
            _local(s.since) if s.since else None,
            s.negative_recent,
            s.negative_before,
            s.trend,
            s.response_rate,
            s.negative_response_rate,
            s.response_hours,
            s.smaller,
            s.bigger,
            s.top_theme or "—",
            "\n".join(conclusions(s)),
        ]
        for s in summaries
    ]
    _table(sheet, 3, SUMMARY_COLUMNS, rows)
    trend_column = 1 + [title for title, _, _ in SUMMARY_COLUMNS].index("Тренд")
    for row, s in enumerate(summaries, start=4):
        if s.trend in TREND_FILLS:
            sheet.cell(row, trend_column).fill = TREND_FILLS[s.trend]
    sheet.freeze_panes = "C4"


def _complaints(sheet: Worksheet, summaries: Sequence[Summary]) -> None:
    _title(
        sheet,
        "Доля жалоб (отзывов на 1–3★ с текстом), где упоминается тема. Один отзыв может попасть в несколько тем",
    )
    columns = [("Тема", 34.0, None), *((_label(s), 18.0, PERCENT) for s in summaries)]
    mentioned = Counter[str]()
    for s in summaries:
        mentioned.update(s.themes)
    rows: list[list[Any]] = []
    for theme in [*(theme for theme, _ in mentioned.most_common()), NO_THEME]:
        counts = [s.unclassified if theme == NO_THEME else s.themes.get(theme, 0) for s in summaries]
        rows.append([theme, *(n / s.worded if s.worded else None for n, s in zip(counts, summaries, strict=True))])
    _table(sheet, 3, columns, rows, autofilter=False)
    totals = (("Жалоб с текстом", [s.worded for s in summaries]), (SILENT, [s.silent for s in summaries]))
    for row, (label, values) in enumerate(totals, start=4 + len(rows)):
        sheet.cell(row, 1, label).font = Font(bold=True)
        for number, value in enumerate(values, start=2):
            sheet.cell(row, number, value).number_format = INTEGER
    sheet.freeze_panes = "B4"


def _ratings(sheet: Worksheet, items: Sequence[Item]) -> None:
    _title(sheet, "Оценки за всё время — по данным WB")
    columns = [
        ("Артикул WB", 12.0, "0"),
        ("Товар", 34.0, None),
        *((f"{stars}★", 8.0, PERCENT) for stars in range(5, 0, -1)),
        ("Оценок всего", 10.0, INTEGER),
    ]
    rows = []
    for feedback, summary in items:
        total = sum(feedback.distribution.values())
        shares = [feedback.distribution.get(stars, 0) / total if total else None for stars in range(5, 0, -1)]
        rows.append([summary.card.nm, summary.card.name, *shares, total])
    _table(sheet, 3, columns, rows)


def _by_month(sheet: Worksheet, items: Sequence[Item], now: datetime) -> None:
    _title(
        sheet,
        "Доля оценок 1–2★ по месяцам. Пусто — меньше 5 отзывов за месяц или WB уже не отдаёт отзывы за тот месяц",
    )
    series = [monthly(feedback, now) for feedback, _ in items]
    labels = [month.label for month in series[0]]
    products = [(_label(summary), 18.0, PERCENT) for _, summary in items]
    shares = [[label, *(months[index].negative for months in series)] for index, label in enumerate(labels)]
    _table(sheet, 3, [("Месяц", 12.0, None), *products], shares, autofilter=False)

    counts_top = 3 + len(labels) + 3
    sheet.cell(counts_top - 1, 1, "Отзывов за месяц").font = Font(bold=True)
    counts = [[label, *(months[index].count for months in series)] for index, label in enumerate(labels)]
    _table(
        sheet,
        counts_top,
        [("Месяц", 12.0, None), *((t, w, "0") for t, w, _ in products)],
        counts,
        autofilter=False,
    )

    chart = LineChart()
    chart.title = "Доля оценок 1–2★"
    chart.height, chart.width = 9, 20  # сантиметры
    chart.y_axis.number_format = PERCENT
    chart.x_axis.delete = False  # иначе новые версии Excel прячут оси
    chart.y_axis.delete = False
    chart.add_data(
        Reference(sheet, min_col=2, max_col=1 + len(items), min_row=3, max_row=3 + len(labels)),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=3 + len(labels)))
    sheet.add_chart(chart, f"{get_column_letter(len(items) + 3)}3")


def _reviews(sheet: Worksheet, items: Sequence[Item], now: datetime, themes: Themes) -> None:
    window = RECENT + BEFORE
    _title(
        sheet,
        f"Отзывы на 1–3★ за {window.days} дней, новые сверху, до {QUOTES_PER_PRODUCT} на товар. "
        "Имена покупателей не сохраняются, контакты в тексте скрыты",
    )
    rows = []
    for feedback, summary in items:
        complaints = sorted(
            (r for r in feedback.reviews if r.rating <= 3 and not r.excluded and r.created >= now - window),
            key=lambda r: r.created,
            reverse=True,
        )
        for r in complaints[:QUOTES_PER_PRODUCT]:
            rows.append([
                summary.card.nm, _short(summary.card.name), r.nm or None, _local(r.created), r.rating,
                ", ".join(sorted(themes_of(r, themes))), ", ".join(r.tags),
                mask_contacts(r.text), mask_contacts(r.cons), "да" if r.answered else "нет",
            ])  # fmt: skip
    _table(sheet, 3, REVIEW_COLUMNS, rows)
    sheet.freeze_panes = "A4"


def _errors(sheet: Worksheet, errors: Sequence[tuple[str, str]]) -> None:
    _table(sheet, 1, (("Товар", 30, None), ("Причина", 100, None)), [list(error) for error in errors])


def build(
    items: Sequence[Item], errors: Sequence[tuple[str, str]], now: datetime, themes: Themes = DEFAULT_THEMES
) -> Workbook:
    workbook = Workbook()
    del workbook["Sheet"]
    summaries = [summary for _, summary in items]
    _summary(workbook.create_sheet("Сводка"), summaries, now)
    _complaints(workbook.create_sheet("Жалобы"), summaries)
    _ratings(workbook.create_sheet("Оценки"), items)
    _by_month(workbook.create_sheet("По месяцам"), items, now)
    _reviews(workbook.create_sheet("Отзывы 1–3★"), items, now, themes)
    if errors:
        _errors(workbook.create_sheet("Не удалось"), errors)
    return workbook


def write_report(
    path: Path,
    items: Sequence[Item],
    errors: Sequence[tuple[str, str]],
    now: datetime,
    themes: Themes = DEFAULT_THEMES,
) -> None:
    workbook = build(items, errors, now, themes)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise ReportError(f"Не удалось записать {path}: файл открыт в Excel? Закройте его и запустите снова.") from exc
